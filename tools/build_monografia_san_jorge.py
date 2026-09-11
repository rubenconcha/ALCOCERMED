from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt, RGBColor


OUTPUT = Path(__file__).resolve().parents[1] / "output" / "documents" / "Monografia_San_Jorge_avance_segunda_revision_2026.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="D9D9D9", size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        el = borders.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), size)
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_width(cell, width_cm):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_cm * 567)))
    tc_w.set(qn("w:type"), "dxa")


def set_fixed_table_layout(table, widths):
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    for index, grid_col in enumerate(table._tbl.tblGrid.gridCol_lst):
        if index < len(widths):
            grid_col.set(qn("w:w"), str(int(widths[index] * 567)))


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:cantSplit")) is None:
        tr_pr.append(OxmlElement("w:cantSplit"))


def set_run_font(run, name="Times New Roman", size=12, bold=None, italic=None, color="000000"):
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_char_1 = OxmlElement("w:fldChar")
    fld_char_1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char_2 = OxmlElement("w:fldChar")
    fld_char_2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char_1, instr_text, fld_char_2])
    set_run_font(run, size=10)


def restart_page_numbering(section, start=1):
    sect_pr = section._sectPr
    pg_num = sect_pr.find(qn("w:pgNumType"))
    if pg_num is None:
        pg_num = OxmlElement("w:pgNumType")
        sect_pr.append(pg_num)
    pg_num.set(qn("w:start"), str(start))


def continue_page_numbering(section):
    """Keep numbering continuous after an orientation-only section break."""
    sect_pr = section._sectPr
    pg_num = sect_pr.find(qn("w:pgNumType"))
    if pg_num is not None:
        sect_pr.remove(pg_num)


def add_toc(paragraph):
    run = paragraph.add_run()
    fld_char = OxmlElement("w:fldChar")
    fld_char.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = ' TOC \\o "1-3" \\h \\z \\u '
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "Actualice el índice en Microsoft Word"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char, instr, separate, placeholder, end])
    set_run_font(run, size=11)


def configure_section(section, landscape=False):
    if landscape:
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width = Mm(297)
        section.page_height = Mm(210)
        section.left_margin = Cm(1.8)
        section.right_margin = Cm(1.8)
        section.top_margin = Cm(1.8)
        section.bottom_margin = Cm(1.8)
    else:
        section.orientation = WD_ORIENT.PORTRAIT
        section.page_width = Mm(210)
        section.page_height = Mm(297)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.5)
    section.header_distance = Cm(1.2)
    section.footer_distance = Cm(1.2)


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    p.add_run(text)
    return p


def add_body(doc, text, bold_lead=None):
    p = doc.add_paragraph(style="Body Text")
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Cm(1.25)
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        set_run_font(r1, bold=True)
        r2 = p.add_run(text[len(bold_lead):])
        set_run_font(r2)
    else:
        r = p.add_run(text)
        set_run_font(r)
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.left_indent = Cm(0.7 + 0.5 * level)
    p.paragraph_format.first_line_indent = Cm(-0.35)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text)
    set_run_font(r)
    return p


def add_numbered_reference(doc, number, text):
    p = doc.add_paragraph(style="Body Text")
    p.paragraph_format.left_indent = Cm(0.8)
    p.paragraph_format.first_line_indent = Cm(-0.8)
    p.paragraph_format.space_after = Pt(6)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p.add_run(f"{number}. {text}")
    set_run_font(r, size=10.5)
    return p


def style_table(table, widths, font_size=8.5, header_fill="1F4E78"):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_fixed_table_layout(table, widths)
    set_table_borders(table)
    for j, width in enumerate(widths):
        table.columns[j].width = Cm(width)
    for i, row in enumerate(table.rows):
        prevent_row_split(row)
        if i == 0:
            set_repeat_table_header(row)
        for j, cell in enumerate(row.cells):
            cell.width = Cm(widths[j])
            set_cell_width(cell, widths[j])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if i == 0:
                set_cell_shading(cell, header_fill)
            elif i % 2 == 0:
                set_cell_shading(cell, "EAF2F8")
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.0
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                for run in p.runs:
                    set_run_font(run, size=font_size, bold=(i == 0), color=("FFFFFF" if i == 0 else "000000"))


def add_label_line(doc, label, text=""):
    p = doc.add_paragraph(style="Body Text")
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(label)
    set_run_font(r, bold=True)
    r2 = p.add_run(text)
    set_run_font(r2)
    return p


def build_document():
    doc = Document()
    configure_section(doc.sections[0], landscape=False)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    normal.font.size = Pt(12)

    body = styles["Body Text"]
    body.font.name = "Times New Roman"
    body._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    body._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    body.font.size = Pt(12)
    body.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    body.paragraph_format.space_after = Pt(6)

    for level, size in ((1, 14), (2, 12), (3, 12)):
        style = styles[f"Heading {level}"]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(12 if level == 1 else 8)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True

    if "Document Subtitle" not in styles:
        subtitle = styles.add_style("Document Subtitle", WD_STYLE_TYPE.PARAGRAPH)
    else:
        subtitle = styles["Document Subtitle"]
    subtitle.font.name = "Times New Roman"
    subtitle.font.size = Pt(13)
    subtitle.font.color.rgb = RGBColor(0, 0, 0)
    subtitle.paragraph_format.space_after = Pt(8)

    # Cover
    cover_lines = [
        ("UNIVERSIDAD MAYOR REAL Y PONTIFICIA DE SAN FRANCISCO XAVIER DE CHUQUISACA", 13, True),
        ("RESIDENCIA MÉDICA UNIDAD DE POSTGRADO", 12, True),
        ("COMITÉ REGIONAL DE INTEGRACIÓN DOCENTE ASISTENCIAL", 11, True),
        ("INVESTIGACIÓN E INTERACCIÓN COMUNITARIA", 11, True),
        ("SUBCOMISIÓN DE POSTGRADO Y RESIDENCIA MÉDICA", 11, True),
    ]
    for idx, (line, size, bold) in enumerate(cover_lines):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(4)
        if idx == 0:
            p.paragraph_format.space_before = Pt(8)
        r = p.add_run(line)
        set_run_font(r, size=size, bold=bold)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(54)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("EXPOSICIÓN ERGONÓMICA LABORAL Y PATRONES DE USO DE ANTIINFLAMATORIOS NO ESTEROIDEOS EN PACIENTES DE 30 A 70 AÑOS CON DOLOR OSTEOARTICULAR ATENDIDOS EN EL CENTRO DE SALUD INTEGRAL SAN JORGE ZUDÁÑEZ CHUQUISACA JULIO A SEPTIEMBRE DE 2026")
    set_run_font(r, size=15, bold=True)

    p = doc.add_paragraph(style="Document Subtitle")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(18)
    r = p.add_run("Avance para segunda revisión")
    set_run_font(r, size=13, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(72)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Autor  Dr. Ruben Concha Cayola")
    set_run_font(r, size=12, bold=True)
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p2.add_run("Médico residente SAFCI")
    set_run_font(r, size=12)
    p3 = doc.add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p3.paragraph_format.space_before = Pt(80)
    r = p3.add_run("Chuquisaca Bolivia")
    set_run_font(r, size=12, bold=True)
    p4 = doc.add_paragraph()
    p4.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p4.add_run("2026")
    set_run_font(r, size=12, bold=True)

    # New section after cover; page numbering begins here.
    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    configure_section(sec, landscape=False)
    sec.footer.is_linked_to_previous = False
    add_page_number(sec.footer.paragraphs[0])
    restart_page_numbering(sec, 1)

    add_heading(doc, "ÍNDICE DE CONTENIDO", 1)
    toc_p = doc.add_paragraph()
    add_toc(toc_p)

    doc.add_page_break()
    add_heading(doc, "LISTA DE ABREVIATURAS Y SÍMBOLOS", 1)
    abbreviations = [
        ("AINE", "Antiinflamatorio no esteroideo"),
        ("COX", "Ciclooxigenasa"),
        ("CSI", "Centro de Salud Integral"),
        ("DM", "Diabetes mellitus"),
        ("ERC", "Enfermedad renal crónica"),
        ("HTA", "Hipertensión arterial"),
        ("IC 95 %", "Intervalo de confianza del 95 por ciento"),
        ("NRS", "Escala numérica de intensidad del dolor"),
        ("OMS", "Organización Mundial de la Salud"),
        ("OR", "Razón de momios"),
        ("RP", "Razón de prevalencias"),
        ("SAFCI", "Salud Familiar Comunitaria Intercultural"),
        ("SUS", "Sistema Único de Salud"),
        ("TME", "Trastornos musculoesqueléticos"),
    ]
    t = doc.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text = "Abreviatura"
    t.rows[0].cells[1].text = "Significado"
    for abbr, meaning in abbreviations:
        row = t.add_row().cells
        row[0].text = abbr
        row[1].text = meaning
    style_table(t, [3.2, 11.5], font_size=10)

    doc.add_page_break()
    add_heading(doc, "I INTRODUCCIÓN", 1)
    add_heading(doc, "1.1 Antecedentes", 2)
    add_body(doc, "Los trastornos musculoesqueléticos abarcan afecciones de músculos, huesos, articulaciones y tejidos conectivos. Con frecuencia se manifiestan mediante dolor persistente, limitación de la movilidad y disminución de la capacidad para trabajar o participar en actividades familiares y comunitarias. La Organización Mundial de la Salud estima que alrededor de 1,71 mil millones de personas viven con alguna condición musculoesquelética y reconoce a este grupo como el principal contribuyente mundial a la discapacidad [1].")
    add_body(doc, "Una parte de esta carga se relaciona con las exigencias físicas del trabajo. La manipulación de cargas, la aplicación de fuerza, los movimientos repetitivos, la flexión o torsión del tronco, las posturas sostenidas y la vibración pueden superar la capacidad de adaptación del sistema musculoesquelético cuando se combinan con alta intensidad, frecuencia o duración [2-4]. Las estimaciones conjuntas de la OMS y la Organización Internacional del Trabajo atribuyeron 12,3 millones de años de vida ajustados por discapacidad a factores ergonómicos ocupacionales en 2016; el dolor de espalda y cuello representó una fracción importante de esa carga [2].")
    add_body(doc, "La evidencia epidemiológica respalda la relación entre tareas específicas y dolor. En trabajadores agrícolas latinos, la flexión del tronco por más de 30 horas semanales se asoció con dolor crónico de cadera, mientras que arrodillarse o gatear durante jornadas prolongadas se relacionó con dolor de espalda y rodilla en mujeres [5]. En trabajadores de mercado de Colombia, las jornadas de ocho a once horas y determinadas actividades de venta se asociaron con mayor prevalencia de síntomas musculoesqueléticos [6]. Estos hallazgos son pertinentes para contextos donde predominan actividades manuales, agrícolas, comerciales y de cuidado no remunerado.")
    add_body(doc, "El dolor osteoarticular suele motivar el uso de antiinflamatorios no esteroideos. Los AINEs ofrecen analgesia y reducción de la inflamación, pero su beneficio depende de la indicación, la vía, la dosis, la duración y el perfil clínico del paciente. Su empleo sin valoración profesional, el uso simultáneo de dos AINEs, la prolongación del tratamiento o la presencia de comorbilidades pueden incrementar el riesgo gastrointestinal, renal y cardiovascular [7-10]. La OMS considera irracional el uso de medicamentos cuando no responde a las necesidades clínicas, emplea dosis o tiempos inadecuados o se realiza mediante automedicación inapropiada [7].")
    add_body(doc, "En América Latina se han descrito niveles relevantes de automedicación. Un estudio realizado en farmacias comunitarias del sur de Chile encontró que los AINEs fueron el grupo más solicitado entre quienes se automedicaban [11]. Un análisis de una encuesta nacional peruana informó una prevalencia de automedicación con AINEs de 68,2 % entre los participantes incluidos y mostró asociación con la compra de productos de marca y de venta sin receta [12]. Aunque estas cifras no pueden extrapolarse directamente a Zudáñez, muestran la necesidad de medir el problema con datos locales.")
    add_body(doc, "El Centro de Salud Integral San Jorge se encuentra en el municipio de Zudáñez, departamento de Chuquisaca, y forma parte del primer nivel de atención descrito en documentos institucionales [13,14]. Su cercanía con la población permite integrar la evaluación clínica del dolor con la identificación de exposiciones laborales y prácticas de uso de medicamentos, elementos necesarios para orientar educación, prevención ergonómica y prescripción segura desde el enfoque SAFCI.")

    add_heading(doc, "1.2 Epidemiología", 2)
    add_body(doc, "Las condiciones musculoesqueléticas afectan a personas de todas las edades, pero su frecuencia aumenta con el envejecimiento. La lumbalgia constituye el componente de mayor carga, seguida por fracturas, osteoartritis, dolor cervical y otras afecciones del aparato locomotor [1]. Además de la limitación funcional, estas condiciones generan consultas repetidas, ausentismo, retiro prematuro del trabajo y costos para las familias y los servicios de salud.")
    add_body(doc, "Desde la salud ocupacional, la exposición no depende únicamente del nombre del oficio. Dos personas con la misma ocupación pueden realizar tareas distintas y acumular cargas diferentes. Por ello, la medición epidemiológica debe registrar la fuerza aplicada, la postura, la repetición, la vibración y el tiempo efectivo de exposición. El Instituto Nacional para la Seguridad y Salud Ocupacional de Estados Unidos destaca que el riesgo se modifica por la intensidad, la frecuencia y la duración, y aumenta cuando varios factores se presentan simultáneamente [4].")
    add_body(doc, "La automedicación con analgésicos y AINEs es un problema farmacoepidemiológico porque puede ocultar la causa del dolor, retrasar la consulta y exponer a eventos adversos evitables. El riesgo no es uniforme: aumenta con la edad, antecedentes de úlcera o hemorragia digestiva, enfermedad renal, hipertensión, insuficiencia cardiaca, enfermedad cardiovascular y uso concomitante de anticoagulantes, antiagregantes, corticoides, diuréticos o medicamentos que actúan sobre el sistema renina angiotensina [8-10]. En personas de 30 a 70 años atendidas en el primer nivel, estas condiciones pueden coexistir y deben formar parte de la evaluación.")
    add_body(doc, "No se identificó una estimación publicada específica que integre exposición ergonómica laboral y patrones de uso de AINEs en pacientes con dolor osteoarticular del CSI San Jorge. Esta ausencia limita la posibilidad de diseñar mensajes educativos y acciones preventivas adaptadas a las tareas y prácticas farmacológicas de la población atendida.")

    add_heading(doc, "1.3 Planteamiento del problema", 2)
    add_body(doc, "El dolor osteoarticular atendido en el primer nivel puede ser la expresión de una enfermedad degenerativa, una lesión previa o una sobrecarga acumulada durante el trabajo. Cuando el manejo se concentra únicamente en aliviar el síntoma, sin caracterizar las exigencias físicas que lo preceden, persiste la exposición y aumenta la probabilidad de recurrencia. De manera paralela, el acceso cotidiano a analgésicos puede favorecer que el paciente repita tratamientos previos, modifique dosis o combine medicamentos sin reconocer contraindicaciones.")
    add_body(doc, "En el CSI San Jorge no se dispone de una caracterización sistemática que muestre qué exposiciones ergonómicas predominan entre los pacientes con dolor osteoarticular, qué AINEs utilizan, quién los indica, durante cuánto tiempo los consumen y qué proporción presenta prácticas potencialmente inseguras. Tampoco se conoce si una mayor carga ergonómica se acompaña de mayor frecuencia de uso, automedicación o patrones de riesgo. Esta brecha de información impide priorizar intervenciones concretas en consulta, visitas familiares y actividades comunitarias.")
    add_body(doc, "Pregunta de investigación: ¿Cuál es la relación entre la exposición ergonómica laboral y los patrones de uso de AINEs en pacientes de 30 a 70 años con dolor osteoarticular atendidos en el Centro de Salud Integral San Jorge, Zudáñez, Chuquisaca, durante julio a septiembre de 2026?", bold_lead="Pregunta de investigación:")

    add_heading(doc, "1.3.1 Justificación", 3)
    add_body(doc, "La investigación es pertinente porque aborda dos componentes modificables de un mismo problema asistencial: las exigencias físicas del trabajo y la manera en que se utilizan los AINEs. Identificar tareas con alta demanda permitirá orientar recomendaciones sobre pausas, organización del trabajo, manipulación de cargas y posturas. Describir el consumo permitirá reconocer automedicación, duplicidad terapéutica, duración prolongada sin control y presencia de factores que elevan el riesgo de reacciones adversas.")
    add_body(doc, "La utilidad práctica radica en que la información puede obtenerse mediante entrevista breve y revisión clínica, sin equipos costosos. Los resultados podrán apoyar consejería individual, educación comunitaria, vigilancia de seguridad farmacológica y derivación oportuna de pacientes con signos de alarma. También aportarán una línea de base para posteriores intervenciones ergonómicas y para mejorar el registro del dolor musculoesquelético en el establecimiento.")
    add_body(doc, "Desde el enfoque SAFCI, el estudio integra la ocupación, el entorno familiar y comunitario, las prácticas de autocuidado y las condiciones clínicas del paciente. El beneficio esperado es institucional y colectivo; no se pretende probar que la exposición laboral cause por sí sola el consumo de AINEs, sino estimar su coexistencia y asociación en un corte temporal para orientar decisiones preventivas.")

    add_heading(doc, "1.3.2 Objetivos", 3)
    add_heading(doc, "Objetivo general", 3)
    add_body(doc, "Determinar la relación entre la exposición ergonómica laboral y los patrones de uso de antiinflamatorios no esteroideos en pacientes de 30 a 70 años con dolor osteoarticular atendidos en el Centro de Salud Integral San Jorge, Zudáñez, Chuquisaca, durante julio a septiembre de 2026.")
    add_heading(doc, "Objetivos específicos", 3)
    for item in [
        "Describir las características sociodemográficas, laborales y clínicas de los pacientes incluidos.",
        "Caracterizar la frecuencia, intensidad y duración de las principales exposiciones ergonómicas laborales relacionadas con fuerza, manipulación de cargas, posturas, repetición y vibración.",
        "Identificar los AINEs utilizados, su vía de administración, frecuencia, duración, fuente de indicación y motivo de uso.",
        "Estimar la proporción de automedicación, uso frecuente y patrones potencialmente inseguros de AINEs.",
        "Analizar la asociación entre el nivel de exposición ergonómica laboral y la frecuencia de uso, la automedicación y los patrones potencialmente inseguros de AINEs, considerando edad, sexo, intensidad del dolor y comorbilidades.",
        "Proponer acciones educativas y preventivas aplicables al primer nivel de atención según los hallazgos del estudio.",
    ]:
        add_bullet(doc, item)

    doc.add_page_break()
    add_heading(doc, "II MARCO TEÓRICO", 1)
    add_heading(doc, "2.1 Dolor osteoarticular y trastornos musculoesqueléticos", 2)
    add_body(doc, "El término trastorno musculoesquelético comprende alteraciones que afectan músculos, tendones, ligamentos, articulaciones, huesos y estructuras relacionadas. El dolor osteoarticular utilizado en este estudio se refiere al dolor localizado en columna o articulaciones de las extremidades que constituye motivo de consulta o se encuentra registrado como problema clínico durante la atención. Esta definición operativa no reemplaza el diagnóstico etiológico y permite incluir cuadros frecuentes del primer nivel, como lumbalgia, cervicalgia, gonalgia, hombro doloroso y artrosis periférica [1].")
    add_body(doc, "La experiencia dolorosa es multidimensional. Su intensidad puede cuantificarse con una escala numérica de cero a diez, pero su repercusión también depende de la duración, el sitio, la limitación funcional, el sueño, las exigencias laborales y las expectativas del paciente. Para el análisis se diferenciará dolor de menos de tres meses y dolor persistente de tres meses o más, sin asumir que todo dolor persistente es de origen ocupacional.")

    add_heading(doc, "2.2 Ergonomía y exposición laboral", 2)
    add_body(doc, "La ergonomía estudia la interacción entre las personas, las tareas, las herramientas y el ambiente de trabajo. En salud ocupacional, una exposición ergonómica es una exigencia física o organizacional capaz de aumentar la carga sobre el aparato locomotor. La presencia de una postura o movimiento aislado no determina daño; el riesgo depende de la combinación entre intensidad, frecuencia, duración, recuperación y capacidad individual [3,4].")
    add_body(doc, "La evaluación clínica de pacientes de diversas ocupaciones requiere preguntar por la tarea real y no limitarse al cargo declarado. El trabajo agrícola, el comercio, la construcción, el transporte, la elaboración de alimentos, las labores domésticas y el cuidado de personas pueden compartir demandas de carga, flexión, empuje, agarre o repetición. También deben considerarse el trabajo informal y no remunerado, porque generan exposición física aunque no exista contrato laboral.")

    add_heading(doc, "2.3 Principales factores ergonómicos", 2)
    factors = [
        ("Fuerza y manipulación de cargas. ", "Comprende levantar, bajar, sostener, transportar, empujar o halar objetos o personas. Su efecto depende del peso, la distancia al cuerpo, la altura, la frecuencia y la técnica."),
        ("Posturas forzadas. ", "Incluyen flexión o torsión del tronco, trabajo con brazos elevados, cuclillas, arrodillamiento y desviaciones sostenidas de muñeca o cuello."),
        ("Posturas estáticas. ", "Mantener por tiempo prolongado una posición de pie, sentado o inclinado reduce la recuperación muscular y puede aumentar fatiga y dolor."),
        ("Repetición. ", "La reiteración de ciclos cortos, especialmente combinada con fuerza o mala postura, incrementa la carga acumulada sobre tendones y articulaciones."),
        ("Vibración. ", "Puede transmitirse a mano y brazo mediante herramientas o a todo el cuerpo mediante vehículos y maquinaria."),
        ("Organización del trabajo. ", "Las jornadas largas, el ritmo elevado, la ausencia de pausas y el bajo control sobre la tarea pueden aumentar la exposición efectiva y la tensión muscular."),
    ]
    for lead, tail in factors:
        p = doc.add_paragraph(style="Body Text")
        p.paragraph_format.left_indent = Cm(0.6)
        p.paragraph_format.first_line_indent = Cm(-0.4)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        r = p.add_run(lead)
        set_run_font(r, bold=True)
        r2 = p.add_run(tail)
        set_run_font(r2)

    add_heading(doc, "2.4 Medición de la exposición ergonómica", 2)
    add_body(doc, "Los métodos de evaluación incluyen observación directa, análisis de tareas, mediciones instrumentales y cuestionarios. Herramientas como Quick Exposure Check integran observación y participación del trabajador, mientras que los cuestionarios de carga física permiten estudiar fuerza, carga dinámica y estática, repetición, postura y vibración [15,16]. En un establecimiento de salud, la entrevista estructurada es factible para comparar pacientes de distintos oficios, pero puede presentar sesgo de recuerdo y no sustituye una evaluación ergonómica del puesto.")
    add_body(doc, "El presente estudio empleará un módulo breve adaptado a partir de dominios descritos por NIOSH y por cuestionarios de carga física. Cada exposición se registrará en una escala de frecuencia de cuatro categorías. El puntaje total se analizará principalmente como variable continua; sus terciles se utilizarán únicamente para describir niveles relativos de menor, intermedia y mayor exposición dentro de la muestra. Estas categorías no equivalen a una certificación de riesgo laboral.")

    add_heading(doc, "2.5 Relación entre trabajo y dolor", 2)
    add_body(doc, "La sobrecarga mecánica repetida puede producir fatiga, microlesiones y menor tolerancia tisular. Las posturas forzadas modifican los brazos de palanca y aumentan la demanda muscular; la repetición limita la recuperación; la vibración y el manejo de cargas pueden potenciar la tensión sobre columna y extremidades. La edad, el acondicionamiento físico, lesiones previas, obesidad, enfermedades crónicas y factores psicosociales pueden modificar la respuesta [3-5].")
    add_body(doc, "Debido a que el diseño es transversal y todos los participantes consultarán por dolor, el estudio no estimará incidencia ni demostrará causalidad. La exposición se interpretará como antecedente asociado. También se considerará la posibilidad de causalidad inversa: una persona con dolor puede haber reducido tareas exigentes antes de la entrevista.")

    add_heading(doc, "2.6 Antiinflamatorios no esteroideos", 2)
    add_body(doc, "Los AINEs inhiben las enzimas ciclooxigenasas y disminuyen la síntesis de prostaglandinas. Este mecanismo explica su acción analgésica y antiinflamatoria, pero también parte de sus efectos adversos. Los AINEs no selectivos inhiben COX 1 y COX 2 en diferente grado; los inhibidores selectivos de COX 2 reducen algunos efectos gastrointestinales superiores, aunque no eliminan el riesgo renal o cardiovascular [8-10].")
    add_body(doc, "Entre los principios activos de uso habitual se encuentran ibuprofeno, diclofenaco, naproxeno, ketorolaco y meloxicam. Pueden administrarse por vía oral, tópica o parenteral. Para fines del estudio, paracetamol, corticoides, opioides y productos herbales se registrarán como analgésicos concomitantes, pero no se clasificarán como AINEs.")

    add_heading(doc, "2.7 Patrones de uso de AINEs", 2)
    add_body(doc, "Un patrón de uso se construye con varios atributos: principio activo, vía, dosis declarada, número de tomas, días de consumo, motivo, fuente de indicación, lugar de obtención y uso simultáneo de otros medicamentos. La automedicación se definirá como el uso de un AINE por decisión propia o recomendación no emitida por un profesional habilitado para el episodio actual, incluyendo la repetición de recetas antiguas sin reevaluación.")
    add_body(doc, "El uso frecuente se definirá operacionalmente como consumo de un AINE sistémico durante diez o más días en los últimos treinta días. Este punto de corte permitirá comparar grupos dentro del estudio y no debe interpretarse como umbral universal de toxicidad. La exposición acumulada será más informativa cuando se combine con la dosis, la duración continua y los factores clínicos del paciente.")

    add_heading(doc, "2.8 Uso racional y patrón potencialmente inseguro", 2)
    add_body(doc, "El uso racional implica seleccionar el medicamento apropiado para la necesidad clínica, en una dosis individual adecuada, durante el tiempo necesario y con seguimiento proporcional al riesgo [7]. Las guías para osteoartritis recomiendan priorizar medidas no farmacológicas, considerar AINEs tópicos en articulaciones accesibles y, cuando se requiera un AINE oral, valorar toxicidad gastrointestinal, renal, hepática y cardiovascular, además de edad, embarazo, comorbilidades y tratamientos concomitantes [17].")
    add_body(doc, "En esta investigación se considerará patrón potencialmente inseguro la presencia de al menos una de las siguientes condiciones: uso simultáneo de dos o más AINEs sistémicos; dosis o frecuencia superior a la indicada; uso continuo por siete o más días sin evaluación profesional; automedicación en presencia de antecedente de úlcera o hemorragia digestiva, enfermedad renal, insuficiencia cardiaca, cardiopatía isquémica, hipertensión no controlada o tratamiento anticoagulante; o reaparición de signos de alarma durante el consumo. Esta clasificación servirá para tamizaje y no sustituirá el juicio clínico.")

    add_heading(doc, "2.9 Riesgos gastrointestinales renales y cardiovasculares", 2)
    add_body(doc, "La inhibición de prostaglandinas puede reducir la protección de la mucosa gastrointestinal y favorecer dispepsia, úlcera, sangrado o perforación. El riesgo aumenta con edad avanzada, antecedente ulceroso, dosis altas, combinación de AINEs, anticoagulantes, antiagregantes o corticoides. La gastroprotección debe individualizarse según el riesgo y no convierte en inocuo el tratamiento [9,10,17].")
    add_body(doc, "A nivel renal, los AINEs pueden disminuir la perfusión glomerular, retener sodio y favorecer lesión renal aguda, especialmente en deshidratación, enfermedad renal previa, insuficiencia cardiaca, diabetes, edad avanzada o uso combinado con diuréticos e inhibidores del sistema renina angiotensina [9]. En el sistema cardiovascular pueden elevar la presión arterial, agravar edema e insuficiencia cardiaca y aumentar eventos trombóticos en determinados pacientes, dosis y duraciones [8,10].")

    add_heading(doc, "2.10 Papel del primer nivel de atención", 2)
    add_body(doc, "El primer nivel puede actuar antes de que el dolor se vuelva recurrente o el consumo farmacológico se prolongue. La atención debe explorar la tarea que desencadena o agrava el dolor, brindar recomendaciones factibles, promover actividad física y rehabilitación cuando corresponda, revisar comorbilidades e interacciones y explicar signos de alarma. El abordaje se complementa con visitas familiares, educación grupal y coordinación con empleadores u organizaciones comunitarias cuando resulte viable.")

    add_heading(doc, "2.11 Modelo conceptual del estudio", 2)
    add_body(doc, "El modelo propone que una mayor exposición a fuerza, posturas forzadas, repetición, vibración y jornadas sin recuperación puede asociarse con mayor intensidad o persistencia del dolor. A su vez, el dolor puede aumentar la frecuencia de uso de AINEs. La decisión de automedicarse o mantener un patrón potencialmente inseguro también depende de acceso al servicio, experiencias previas, disponibilidad de medicamentos, conocimientos, edad y comorbilidades. Por ello, la relación principal será ajustada por variables clínicas y sociodemográficas relevantes.")
    add_body(doc, "En términos analíticos, el puntaje de exposición ergonómica constituye la variable independiente principal, mientras que los días de uso, la automedicación, el uso frecuente y el patrón potencialmente inseguro de AINEs son variables de resultado. La intensidad y la duración del dolor pueden funcionar como variables intermedias: una carga física elevada puede agravar el síntoma y este, a su vez, motivar un mayor consumo. La edad, el sexo, la ocupación, la antigüedad laboral y las comorbilidades se tratarán como posibles factores de confusión cuando la distribución de los datos permita el ajuste.")
    add_body(doc, "El modelo también reconoce factores contextuales propios del primer nivel, como la disponibilidad de medicamentos, la distancia o facilidad de acceso, la consulta previa y la información recibida. Estos elementos no se asumirán como causas directas, sino que se describirán para interpretar por qué dos pacientes con exposición y dolor semejantes pueden adoptar conductas farmacológicas diferentes. La lectura final deberá integrar magnitud de la asociación, precisión estadística, plausibilidad clínica y limitaciones del diseño transversal.")

    doc.add_page_break()
    add_heading(doc, "III METODOLOGÍA DE LA INVESTIGACIÓN", 1)
    add_heading(doc, "3.1 Formulación de hipótesis", 2)
    add_body(doc, "Hipótesis de investigación H1: Los pacientes con mayor exposición ergonómica laboral presentan mayor frecuencia de uso, automedicación o patrones potencialmente inseguros de AINEs que los pacientes con menor exposición, durante el periodo de estudio.", bold_lead="Hipótesis de investigación H1:")
    add_body(doc, "Hipótesis nula H0: No existe asociación estadísticamente significativa entre el nivel de exposición ergonómica laboral y los patrones de uso de AINEs en los pacientes incluidos.", bold_lead="Hipótesis nula H0:")

    add_heading(doc, "3.2 Identificación de variables", 2)
    add_label_line(doc, "Variable de exposición principal: ", "puntaje de exposición ergonómica laboral.")
    add_label_line(doc, "Variables de resultado: ", "días de uso de AINEs en los últimos 30 días, automedicación, uso frecuente y patrón potencialmente inseguro.")
    add_label_line(doc, "Variables clínicas intermedias: ", "intensidad, localización y duración del dolor; limitación de actividad.")
    add_label_line(doc, "Variables de control: ", "edad, sexo, procedencia, ocupación, antigüedad laboral, jornada, comorbilidades, medicamentos concomitantes y acceso al servicio.")

    # Landscape operationalization table.
    landscape = doc.add_section(WD_SECTION.NEW_PAGE)
    configure_section(landscape, landscape=True)
    continue_page_numbering(landscape)
    landscape.footer.is_linked_to_previous = True
    add_heading(doc, "3.3 Operacionalización de variables", 2)
    headers = ["Variable", "Definición conceptual", "Definición operacional", "Indicador", "Categorías o valores", "Escala"]
    rows = [
        ["Exposición ergonómica laboral", "Exigencias físicas y organizacionales del trabajo con potencial de sobrecarga musculoesquelética.", "Suma de 18 ítems de frecuencia, puntuados de 0 a 3, sobre fuerza, carga, postura, repetición, vibración y recuperación.", "Puntaje total 0 a 54 y subescalas.", "Continua; terciles relativos para descripción.", "Razón y ordinal"],
        ["Manipulación de cargas", "Levantamiento, transporte, empuje o tracción de objetos o personas.", "Frecuencia declarada y peso habitual de la carga en la actividad principal.", "Días por semana, proporción de jornada y kg.", "Nunca, ocasional, frecuente, casi siempre; <5, 5-9, 10-19, ≥20 kg.", "Ordinal y razón"],
        ["Posturas forzadas", "Posiciones articulares alejadas de la postura neutral.", "Flexión o torsión del tronco, arrodillamiento, cuclillas o brazos elevados durante el trabajo.", "Frecuencia por ítem.", "0 nunca a 3 casi siempre.", "Ordinal"],
        ["Movimientos repetitivos", "Repetición de ciclos de movimiento con recuperación limitada.", "Movimientos de manos, brazos o tronco repetidos durante la actividad principal.", "Frecuencia percibida.", "0 nunca a 3 casi siempre.", "Ordinal"],
        ["Vibración", "Energía mecánica transmitida al cuerpo por herramientas, equipos o vehículos.", "Uso de herramientas vibrátiles o conducción de maquinaria durante el trabajo.", "Presencia y frecuencia.", "No; sí ocasional; sí frecuente.", "Ordinal"],
        ["Uso de AINEs", "Consumo de un fármaco antiinflamatorio no esteroideo.", "Uso por vía oral, intramuscular, intravenosa o tópica para el dolor actual en los últimos 30 días.", "Principio activo, vía y días de uso.", "No; sí. Detalle por medicamento.", "Nominal y razón"],
        ["Uso frecuente de AINE sistémico", "Consumo repetido durante una proporción relevante del periodo.", "AINE oral o parenteral utilizado 10 o más días de los últimos 30.", "Número de días.", "<10 días; ≥10 días.", "Nominal dicotómica"],
        ["Automedicación", "Uso de medicamentos sin indicación profesional actual.", "Decisión propia, recomendación familiar, de conocido o dependiente de farmacia, o reutilización de receta antigua sin evaluación para el episodio actual.", "Fuente de indicación.", "Sí; no.", "Nominal dicotómica"],
        ["Patrón potencialmente inseguro", "Práctica que puede elevar el riesgo de evento adverso o uso inadecuado.", "Al menos un criterio predefinido: duplicidad, exceso de dosis, ≥7 días sin control, automedicación con contraindicación o signos de alarma.", "Número y tipo de criterios.", "Sí; no.", "Nominal dicotómica"],
        ["Intensidad del dolor", "Magnitud percibida del dolor.", "Puntuación de 0 a 10 en escala numérica al momento de la entrevista.", "Puntaje NRS.", "0; 1-3; 4-6; 7-10.", "Ordinal"],
        ["Duración del dolor", "Tiempo transcurrido desde el inicio del episodio o persistencia del síntoma.", "Tiempo declarado del dolor actual.", "Días o meses.", "<3 meses; ≥3 meses.", "Razón y nominal"],
        ["Comorbilidad de riesgo", "Condición clínica que modifica la seguridad de los AINEs.", "Antecedente documentado o declarado de úlcera o sangrado, ERC, HTA, insuficiencia cardiaca, cardiopatía isquémica, enfermedad hepática o diabetes.", "Número y tipo de condiciones.", "Ninguna; una; dos o más.", "Ordinal"],
        ["Edad", "Tiempo de vida desde el nacimiento.", "Años cumplidos al momento de la entrevista.", "Años.", "30 a 70.", "Razón"],
        ["Sexo", "Característica registrada en la atención.", "Dato consignado en registro clínico y confirmado en entrevista.", "Categoría registrada.", "Femenino; masculino; otro registro disponible.", "Nominal"],
        ["Ocupación principal", "Actividad productiva o de cuidado realizada con mayor dedicación.", "Trabajo actual o más reciente según criterio temporal del protocolo.", "Grupo ocupacional y tarea.", "Agricultura, comercio, construcción, transporte, servicios, cuidado o labores del hogar, otros.", "Nominal"],
    ]
    for block_index, (block_title, block_rows) in enumerate([
        ("Variables principales", rows[:9]),
        ("Variables clínicas y de control", rows[9:]),
    ]):
        if block_index:
            doc.add_page_break()
        block_label = doc.add_paragraph()
        block_label.paragraph_format.space_after = Pt(5)
        block_run = block_label.add_run(block_title)
        set_run_font(block_run, size=9.5, bold=True)
        table = doc.add_table(rows=1, cols=len(headers))
        for i, h in enumerate(headers):
            table.rows[0].cells[i].text = h
        for data in block_rows:
            cells = table.add_row().cells
            for i, value in enumerate(data):
                cells[i].text = value
        style_table(table, [3.0, 4.2, 5.8, 3.2, 4.7, 2.4], font_size=7.8)

    note = doc.add_paragraph()
    note.paragraph_format.space_before = Pt(5)
    note.paragraph_format.space_after = Pt(0)
    note_run = note.add_run("Nota: los terciles del puntaje ergonómico serán relativos a la distribución observada y no representan categorías diagnósticas ni certificación del puesto de trabajo.")
    set_run_font(note_run, size=8.5, italic=True)

    portrait = doc.add_section(WD_SECTION.NEW_PAGE)
    configure_section(portrait, landscape=False)
    continue_page_numbering(portrait)
    portrait.footer.is_linked_to_previous = True

    add_heading(doc, "3.4 Tipo y enfoque del estudio", 2)
    add_body(doc, "Estudio cuantitativo, observacional, analítico y de corte transversal. La recolección será prospectiva durante julio a septiembre de 2026. La exposición laboral y los patrones de uso de AINEs se medirán en una misma entrevista; por ello las asociaciones se interpretarán como relaciones de prevalencia y no como prueba de causalidad.")

    add_heading(doc, "3.5 Área y periodo de estudio", 2)
    add_body(doc, "La investigación se realizará en el Centro de Salud Integral San Jorge, municipio de Zudáñez, departamento de Chuquisaca. Se incluirán las atenciones efectuadas entre el 1 de julio y el 30 de septiembre de 2026 en los espacios asistenciales donde se identifiquen pacientes elegibles.")

    add_heading(doc, "3.6 Universo unidad de análisis y muestra", 2)
    add_heading(doc, "3.6.1 Universo", 3)
    add_body(doc, "El universo estará constituido por todos los pacientes de 30 a 70 años con dolor osteoarticular atendidos en el CSI San Jorge durante el periodo de estudio.")
    add_heading(doc, "3.6.2 Unidad de análisis", 3)
    add_body(doc, "La unidad de análisis será cada paciente elegible que otorgue su consentimiento y complete la entrevista estructurada y la verificación clínica mínima.")
    add_heading(doc, "3.6.3 Muestra y muestreo", 3)
    add_body(doc, "Se realizará inclusión censal de casos consecutivos: se invitará a participar a todos los pacientes que cumplan los criterios durante el periodo definido. Este procedimiento evita seleccionar una fracción arbitraria de una población esperada pequeña y permite documentar el total de elegibles, incluidos, rechazos y formularios incompletos. El tamaño final se informará con el diagrama de participantes y se reconocerá la precisión alcanzada mediante intervalos de confianza.")

    add_heading(doc, "3.7 Criterios de inclusión", 2)
    for item in [
        "Edad cumplida entre 30 y 70 años.",
        "Atención en el CSI San Jorge entre julio y septiembre de 2026.",
        "Dolor localizado en columna o articulaciones de extremidades como motivo de consulta o problema clínico activo.",
        "Actividad laboral, productiva o de cuidado no remunerado actual, o actividad principal ejercida durante al menos seis meses en los últimos cinco años.",
        "Capacidad de responder la entrevista en castellano o con apoyo lingüístico institucional autorizado.",
        "Consentimiento informado para participar.",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "3.8 Criterios de exclusión", 2)
    for item in [
        "Dolor debido exclusivamente a fractura reciente, traumatismo mayor, posoperatorio inmediato o urgencia que requiera atención prioritaria.",
        "Compromiso cognitivo, alteración del estado de conciencia o condición clínica que impida una entrevista válida.",
        "Paciente que no pueda identificar su actividad principal o los medicamentos utilizados y no exista una fuente verificable.",
        "Participación previa en el estudio durante otra consulta del mismo periodo.",
        "Formulario con ausencia de la variable principal de exposición o de la información sobre uso de AINEs.",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "3.9 Técnicas e instrumentos de recolección", 2)
    add_body(doc, "Se utilizará entrevista estructurada, revisión dirigida del registro clínico y verificación de envases, recetas o fotografías del medicamento cuando estén disponibles. La ficha se organiza en cinco módulos: datos sociodemográficos; actividad laboral; exposición ergonómica; características del dolor; y uso y seguridad de AINEs.")
    add_body(doc, "El módulo ergonómico contiene 18 ítems con respuesta de 0 nunca o casi nunca, 1 algunas veces, 2 frecuentemente y 3 casi siempre o durante la mayor parte de la jornada. Los ítems representan dos dominios: trabajo físico pesado y manipulación de cargas; y posturas prolongadas o movimientos repetitivos. Se añadirán preguntas sobre jornada, pausas, años en la actividad y vibración. El instrumento es una adaptación para este estudio y no se presentará como evaluación certificada del puesto [4,15,16].")
    add_body(doc, "El módulo farmacológico registrará principio activo, nombre comercial si se recuerda, vía, dosis por toma, número de tomas diarias, días de uso, fuente de indicación, lugar de obtención, combinación de productos y eventos adversos. Cuando el participante no conozca la dosis, el dato se consignará como desconocido y no se inferirá.")

    add_heading(doc, "3.10 Procedimiento", 2)
    procedure = [
        "Solicitar autorización escrita a la dirección del establecimiento y aprobación académica del protocolo.",
        "Capacitar al investigador o encuestador en elegibilidad, consentimiento, entrevista neutral, identificación de principios activos y criterios de derivación.",
        "Realizar una prueba piloto con participantes no incluidos en el análisis definitivo y ajustar redacción, orden y duración.",
        "Identificar pacientes elegibles después de la evaluación clínica inicial, sin interferir con la atención.",
        "Explicar el estudio, obtener consentimiento y asignar un código único sin registrar nombres en la base analítica.",
        "Aplicar la entrevista en un espacio con privacidad; contrastar datos clínicos y medicamentos cuando exista documentación disponible.",
        "Revisar el formulario el mismo día y corregir únicamente omisiones que puedan aclararse con el participante, sin completar datos por suposición.",
        "Ingresar los datos en una base protegida y realizar doble verificación de una muestra de registros.",
        "Derivar al profesional tratante si se identifica un signo de alarma o un patrón de uso que requiera evaluación clínica inmediata.",
    ]
    for item in procedure:
        add_bullet(doc, item)

    add_heading(doc, "3.11 Plan de análisis", 2)
    add_body(doc, "Se describirán las variables categóricas mediante frecuencias absolutas y porcentajes con intervalos de confianza del 95 %. Las variables cuantitativas se resumirán con media y desviación estándar cuando tengan distribución aproximadamente normal, o con mediana y rango intercuartílico cuando presenten asimetría. Se informará la cantidad de datos faltantes para cada variable.")
    add_body(doc, "El puntaje ergonómico total y sus dominios se analizarán como variables continuas. Para presentación descriptiva se podrán dividir en terciles de la distribución observada, denominados menor, intermedia y mayor exposición relativa. Se comparará la prevalencia de automedicación, uso frecuente y patrón potencialmente inseguro entre categorías mediante prueba de chi cuadrado o prueba exacta de Fisher. Para variables continuas se empleará t de Student o U de Mann Whitney según los supuestos.")
    add_body(doc, "Si el número de participantes y eventos lo permite, se estimarán razones de prevalencia crudas y ajustadas mediante regresión de Poisson con varianza robusta. El modelo incluirá edad, sexo, intensidad del dolor y comorbilidad de riesgo, y limitará el número de covariables para evitar sobreajuste. Se usarán intervalos de confianza del 95 % y un nivel de significación de 0,05. Si la muestra es insuficiente, el análisis multivariable se omitirá y los resultados se presentarán como exploratorios.")

    add_heading(doc, "3.12 Validación y control de calidad", 2)
    add_body(doc, "La validez de contenido será evaluada por al menos tres profesionales con experiencia en medicina familiar o SAFCI, farmacología y salud ocupacional o metodología de investigación. Cada ítem será revisado por claridad, pertinencia y coherencia con los objetivos. La prueba piloto valorará comprensión, tiempo de aplicación, opciones de respuesta y términos locales utilizados para tareas y medicamentos.")
    add_body(doc, "Se elaborará un manual breve de codificación. Las respuestas desconocidas tendrán un código diferenciado y no se convertirán en respuestas negativas. Se comprobarán rangos, duplicados y congruencia entre número de tomas, días de uso y vía. La consistencia interna del módulo ergonómico se explorará con alfa de Cronbach, reconociendo que este indicador no demuestra por sí solo validez del instrumento.")

    add_heading(doc, "3.13 Sesgos y limitaciones", 2)
    add_body(doc, "El muestreo consecutivo puede limitar la generalización a personas que no consultan o acuden a otros establecimientos. La exposición y el consumo se basan parcialmente en memoria y pueden presentar error de clasificación. El dolor puede influir en la percepción de la carga laboral, y algunos pacientes pueden haber modificado su actividad por el mismo síntoma. La verificación de recetas y envases reducirá, pero no eliminará, el error en la identificación de AINEs. El diseño transversal no permite establecer temporalidad causal.")

    add_heading(doc, "3.14 Consideraciones éticas", 2)
    add_body(doc, "El estudio se realizará conforme a los principios de respeto, beneficio, no maleficencia, justicia y confidencialidad establecidos en la Declaración de Helsinki vigente [18]. La participación será voluntaria y no modificará el derecho a la atención. Se solicitará consentimiento informado antes de la entrevista y el participante podrá retirarse sin consecuencia asistencial.")
    add_body(doc, "La base utilizará códigos y no contendrá nombres, números de documento ni teléfonos. Los formularios físicos permanecerán bajo resguardo y los archivos digitales estarán protegidos con contraseña, accesibles únicamente al equipo autorizado. Los resultados se presentarán de forma agregada. Debido a que la investigación consiste en entrevista y revisión clínica, se considera de riesgo mínimo; sin embargo, cualquier signo de alarma o uso potencialmente peligroso identificado será comunicado al profesional tratante para evaluación.")

    doc.add_page_break()
    add_heading(doc, "REFERENCIAS BIBLIOGRÁFICAS", 1)
    refs = [
        "World Health Organization. Musculoskeletal health [Internet]. Geneva: WHO; 2022 [citado 10 sep 2026]. Disponible en: https://www.who.int/news-room/fact-sheets/detail/musculoskeletal-conditions",
        "World Health Organization; International Labour Organization. Total WHO ILO Joint Estimates of the Work related Burden of Disease and Injury 2000 to 2016 [Internet]. Geneva: WHO; 2023 [citado 10 sep 2026]. Disponible en: https://www.who.int/news-room/questions-and-answers/item/total-who-ilo-joint-estimates-of-the-work-related-burden-of-disease-and-injury--2000-2016/",
        "World Health Organization. Preventing musculoskeletal disorders in the workplace. Protecting Workers Health Series No. 5. Geneva: WHO; 2003.",
        "National Institute for Occupational Safety and Health. Step 1 Identify risk factors [Internet]. Atlanta: CDC; 2024 [citado 10 sep 2026]. Disponible en: https://www.cdc.gov/niosh/ergonomics/ergo-programs/risk-factors.html",
        "Xiao H, McCurdy SA, Stoecklin Marois MT, Li CS, Schenker MB. Agricultural work and chronic musculoskeletal pain among Latino farm workers the MICASA study. Am J Ind Med. 2013;56(2):216-225. doi:10.1002/ajim.22118.",
        "Garzón Duque MO, Bonbón Velez MC, Toro Rivera JA, Agudelo Aguilar I. Sociodemographic and labor conditions and the presence of musculoskeletal symptoms in workers in a market in a Colombian municipality. Rev Bras Med Trab. 2022;20(2):298-310. doi:10.47626/1679-4435-2022-687.",
        "World Health Organization. Promoting rational use of medicines [Internet]. Geneva: WHO [citado 10 sep 2026]. Disponible en: https://www.who.int/activities/promoting-rational-use-of-medicines/",
        "Minhas D, Nidamarthy S, Gabriel S, Shmagel A. Recommendations for the use of nonsteroidal anti inflammatory drugs and cardiovascular disease risk decades later any new lessons learned. Rheum Dis Clin North Am. 2023;49(1):179-191. doi:10.1016/j.rdc.2022.09.007.",
        "Baker M, Perazella MA. NSAIDs in CKD are they safe. Am J Kidney Dis. 2020;76(4):546-557. doi:10.1053/j.ajkd.2020.03.023.",
        "Sostres C, Gargallo CJ, Lanas A. Gastrointestinal and cardiovascular adverse events associated with NSAIDs. Expert Opin Drug Saf. 2021;20(10):1197-1210. doi:10.1080/14740338.2021.1952988.",
        "Fuentes Albarrán K, Villa Zapata L. Analysis and quantification of self medication patterns of customers in community pharmacies in southern Chile. Pharm World Sci. 2008;30(6):863-868. doi:10.1007/s11096-008-9241-4.",
        "Benites Meza JK, Pinedo Castillo L, Cabanillas Lazo M, et al. Self medication with NSAIDs and purchase of branded and over the counter medicines analysis of a national survey in Peru. J Public Health Res. 2025;14(1):22799036251319154. doi:10.1177/22799036251319154.",
        "Cámara de Diputados del Estado Plurinacional de Bolivia. Legisladores fiscalizan construcción de centro de salud en municipio de Zudáñez [Internet]. La Paz: Cámara de Diputados; 2024 [citado 10 sep 2026]. Disponible en: https://diputados.gob.bo/noticias/legisladores-fiscalizan-construccion-de-centro-de-salud-en-municipio-de-zudanez/",
        "Defensoría del Pueblo. Situación de los Centros de Salud Integrales de primer nivel en el marco del Sistema Nacional de Salud. La Paz: Defensoría del Pueblo; 2018.",
        "David G, Woods V, Li G, Buckle P. The development of the Quick Exposure Check for assessing exposure to risk factors for work related musculoskeletal disorders. Appl Ergon. 2008;39(1):57-69. doi:10.1016/j.apergo.2007.03.002.",
        "Bot SDM, Terwee CB, van der Windt DAWM, Feleus A, Bierma Zeinstra SMA, Knol DL, et al. Internal consistency and validity of a new physical workload questionnaire. Occup Environ Med. 2004;61(12):980-986. doi:10.1136/oem.2003.011213.",
        "National Institute for Health and Care Excellence. Osteoarthritis in over 16s diagnosis and management. NICE guideline NG226 [Internet]. London: NICE; 2022 [citado 10 sep 2026]. Disponible en: https://www.nice.org.uk/guidance/ng226/chapter/Recommendations",
        "World Medical Association. WMA Declaration of Helsinki ethical principles for medical research involving human participants [Internet]. Ferney Voltaire: WMA; 2024 [citado 10 sep 2026]. Disponible en: https://www.wma.net/policies-post/wma-declaration-of-helsinki/",
    ]
    for i, ref in enumerate(refs, 1):
        add_numbered_reference(doc, i, ref)

    doc.add_page_break()
    add_heading(doc, "ANEXOS", 1)
    add_heading(doc, "Anexo 1 Ficha de recolección de datos", 2)
    add_label_line(doc, "Código: ", "______________     Fecha: ____ / ____ / 2026     Encuestador: ____________")
    add_label_line(doc, "Criterios confirmados: ", "[ ] Edad 30 a 70   [ ] Dolor osteoarticular   [ ] Actividad laboral   [ ] Consentimiento")

    add_heading(doc, "A Datos sociodemográficos y clínicos", 3)
    demo = doc.add_table(rows=0, cols=2)
    demo_data = [
        ["Edad en años", "________"],
        ["Sexo registrado", "[ ] Femenino   [ ] Masculino   [ ] Otro registro"],
        ["Residencia", "[ ] Urbana   [ ] Rural   Comunidad: __________"],
        ["Diagnóstico o motivo clínico", "________________________________________"],
        ["Localización principal", "[ ] Cuello [ ] Hombro [ ] Codo [ ] Mano [ ] Espalda [ ] Cadera [ ] Rodilla [ ] Pie [ ] Otra"],
        ["Duración del dolor", "_____ días / _____ meses   [ ] Menos de 3 meses   [ ] 3 meses o más"],
        ["Intensidad actual", "0  1  2  3  4  5  6  7  8  9  10"],
        ["Limitación de actividad", "[ ] Ninguna [ ] Leve [ ] Moderada [ ] Severa"],
    ]
    for data in demo_data:
        cells = demo.add_row().cells
        cells[0].text, cells[1].text = data
    style_table(demo, [5.0, 9.7], font_size=9, header_fill="D9EAF7")
    for p in demo.rows[0].cells[0].paragraphs + demo.rows[0].cells[1].paragraphs:
        for r in p.runs:
            set_run_font(r, size=9, bold=False, color="000000")

    add_heading(doc, "B Actividad laboral principal", 3)
    labor = doc.add_table(rows=0, cols=2)
    labor_data = [
        ["Actividad u oficio", "________________________________________"],
        ["Tarea principal", "________________________________________"],
        ["Condición", "[ ] Dependiente [ ] Cuenta propia [ ] Familiar no remunerado [ ] Cuidado o hogar [ ] Otro"],
        ["Antigüedad", "_____ años _____ meses"],
        ["Jornada", "_____ horas por día   _____ días por semana"],
        ["Pausas de al menos 5 minutos", "[ ] No [ ] Sí   Número aproximado por jornada: ____"],
        ["Actividad actual", "[ ] Sí [ ] No, corresponde a la más reciente de los últimos 5 años"],
    ]
    for data in labor_data:
        cells = labor.add_row().cells
        cells[0].text, cells[1].text = data
    style_table(labor, [5.0, 9.7], font_size=9, header_fill="D9EAF7")
    for row in labor.rows:
        for cell in row.cells:
            set_cell_shading(cell, "FFFFFF")
            for p in cell.paragraphs:
                for r in p.runs:
                    set_run_font(r, size=9, bold=False, color="000000")

    add_heading(doc, "C Módulo de exposición ergonómica", 3)
    add_body(doc, "Marque una opción según la frecuencia habitual en la actividad principal: 0 nunca o casi nunca; 1 algunas veces; 2 frecuentemente; 3 casi siempre o durante la mayor parte de la jornada.")
    ergo_items = [
        "Realiza trabajo físicamente pesado.",
        "Levanta o baja cargas de 5 kg o más.",
        "Levanta o transporta cargas de 10 kg o más.",
        "Empuja o hala objetos, equipos, animales o personas.",
        "Aplica fuerza intensa con manos o brazos.",
        "Trabaja con el tronco inclinado hacia adelante.",
        "Trabaja con el tronco girado o en posición asimétrica.",
        "Permanece en cuclillas, arrodillado o agachado.",
        "Trabaja con uno o ambos brazos por encima del hombro.",
        "Permanece de pie sin cambiar de posición por periodos prolongados.",
        "Permanece sentado por periodos prolongados.",
        "Realiza movimientos repetidos de manos o muñecas.",
        "Realiza movimientos repetidos de brazos u hombros.",
        "Realiza movimientos repetidos del tronco.",
        "Utiliza herramientas que vibran.",
        "Conduce vehículos o maquinaria con vibración de cuerpo entero.",
        "Trabaja a ritmo elevado o con pocas oportunidades de recuperación.",
        "Finaliza la jornada con fatiga física marcada.",
    ]
    et = doc.add_table(rows=1, cols=6)
    for i, h in enumerate(["N", "Exposición", "0", "1", "2", "3"]):
        et.rows[0].cells[i].text = h
    for idx, item in enumerate(ergo_items, 1):
        cells = et.add_row().cells
        values = [str(idx), item, "[ ]", "[ ]", "[ ]", "[ ]"]
        for j, value in enumerate(values):
            cells[j].text = value
            if j != 1:
                cells[j].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_table(et, [0.9, 10.5, 0.8, 0.8, 0.8, 0.8], font_size=8.2)
    add_label_line(doc, "Peso máximo habitual de carga: ", "[ ] <5 kg [ ] 5 a 9 kg [ ] 10 a 19 kg [ ] 20 kg o más [ ] No sabe")
    add_label_line(doc, "Puntaje total: ", "______ / 54")

    doc.add_page_break()
    add_heading(doc, "D Uso de medicamentos para el dolor en los últimos 30 días", 3)
    add_label_line(doc, "¿Usó algún medicamento para este dolor? ", "[ ] No [ ] Sí")
    meds = doc.add_table(rows=1, cols=7)
    med_headers = ["Medicamento o marca", "Vía", "Dosis por toma", "Tomas por día", "Días usados", "Quién indicó", "Continúa"]
    for i, h in enumerate(med_headers):
        meds.rows[0].cells[i].text = h
    for _ in range(4):
        cells = meds.add_row().cells
        for i in range(7):
            cells[i].text = ""
    style_table(meds, [3.2, 1.8, 2.2, 2.0, 1.8, 3.0, 1.5], font_size=8)
    med_questions = [
        "Fuente de indicación: [ ] Médico [ ] Odontólogo [ ] Enfermería según protocolo [ ] Dependiente de farmacia [ ] Familiar o conocido [ ] Decisión propia [ ] Receta antigua",
        "Lugar de obtención: [ ] Centro de salud [ ] Farmacia [ ] Tienda [ ] Familiar [ ] Remanente en casa [ ] Otro",
        "¿Usó dos AINEs al mismo tiempo? [ ] No [ ] Sí   Cuáles: ______________________________",
        "¿Aumentó la dosis o frecuencia indicada? [ ] No [ ] Sí [ ] No recibió indicación",
        "¿Usó un AINE sistémico durante 10 o más días en el último mes? [ ] No [ ] Sí",
        "¿Usó un AINE durante 7 o más días seguidos sin evaluación profesional? [ ] No [ ] Sí",
        "¿Recibió explicación sobre duración, riesgos o signos de alarma? [ ] No [ ] Sí",
    ]
    for q in med_questions:
        add_label_line(doc, "", q)

    add_heading(doc, "E Riesgos y eventos durante el uso", 3)
    risks = [
        "Antecedente de úlcera o sangrado digestivo",
        "Enfermedad renal o creatinina elevada",
        "Hipertensión arterial",
        "Insuficiencia cardiaca, infarto o accidente cerebrovascular",
        "Enfermedad hepática",
        "Diabetes mellitus",
        "Uso de anticoagulante o antiagregante",
        "Uso de corticoide",
        "Uso simultáneo de diurético e inhibidor de la enzima convertidora o antagonista de angiotensina",
    ]
    rt = doc.add_table(rows=1, cols=4)
    for i, h in enumerate(["Condición", "No", "Sí", "No sabe"]):
        rt.rows[0].cells[i].text = h
    for item in risks:
        cells = rt.add_row().cells
        for i, value in enumerate([item, "[ ]", "[ ]", "[ ]"]):
            cells[i].text = value
            if i > 0:
                cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_table(rt, [10.5, 1.4, 1.4, 1.7], font_size=8.5)
    add_label_line(doc, "Síntomas durante el uso: ", "[ ] Dolor o ardor de estómago [ ] Heces negras o vómito con sangre [ ] Menor orina [ ] Hinchazón [ ] Falta de aire [ ] Dolor de pecho [ ] Ninguno")
    add_label_line(doc, "Conducta si corresponde: ", "[ ] Consejería [ ] Evaluación clínica inmediata [ ] Referencia [ ] Sin conducta adicional")
    add_label_line(doc, "Observaciones: ", "________________________________________________________________________")
    for _ in range(6):
        add_label_line(doc, "", "________________________________________________________________________")

    doc.add_page_break()
    add_heading(doc, "Anexo 2 Consentimiento informado para entrevista", 2)
    add_body(doc, "Se invita a participar en un estudio sobre condiciones del trabajo, dolor osteoarticular y uso de antiinflamatorios. La participación consiste en responder una entrevista de aproximadamente 12 a 15 minutos y permitir la revisión de información clínica relacionada con el dolor y los medicamentos utilizados. No se administrará ningún tratamiento como parte de la investigación.")
    add_body(doc, "La participación es voluntaria. Puede negarse o retirarse en cualquier momento sin afectar la atención que recibe. Algunas preguntas podrían causar incomodidad al recordar dolor o prácticas de automedicación; puede omitir una respuesta o solicitar una pausa. Si durante la entrevista se identifica un signo de alarma o una práctica que requiera valoración, se comunicará al profesional tratante.")
    add_body(doc, "La información será codificada y analizada de manera agrupada. No se publicarán nombres ni datos que permitan identificarle. No se ofrece pago por participar. El beneficio directo no está garantizado, pero los resultados pueden orientar actividades de prevención ergonómica y uso seguro de medicamentos en el establecimiento.")
    add_label_line(doc, "Declaración del participante: ", "He recibido una explicación comprensible, pude realizar preguntas y acepto participar voluntariamente.")
    add_label_line(doc, "Nombre o código del participante: ", "____________________________________________")
    add_label_line(doc, "Firma o huella: ", "________________________   Fecha: ____ / ____ / 2026")
    add_label_line(doc, "Responsable del consentimiento: ", "____________________________________________")
    add_label_line(doc, "Firma: ", "________________________")

    doc.add_page_break()
    add_heading(doc, "Anexo 3 Solicitud de autorización institucional", 2)
    add_label_line(doc, "Zudáñez: ", "____ de ______________ de 2026")
    add_label_line(doc, "A: ", "Dirección del Centro de Salud Integral San Jorge")
    add_label_line(doc, "Ref: ", "Solicitud de autorización para investigación")
    add_body(doc, "De mi consideración: Por medio de la presente solicito autorización para realizar la investigación titulada Exposición ergonómica laboral y patrones de uso de antiinflamatorios no esteroideos en pacientes de 30 a 70 años con dolor osteoarticular atendidos en el Centro de Salud Integral San Jorge, Zudáñez, Chuquisaca, julio a septiembre de 2026.")
    add_body(doc, "El estudio consistirá en entrevistas estructuradas y revisión dirigida de información clínica, previa aceptación voluntaria de los participantes. La recolección no interferirá con la atención y los datos se registrarán mediante códigos, sin identificadores personales en la base de análisis. Cualquier hallazgo que requiera atención será comunicado al profesional responsable.")
    add_body(doc, "Agradeciendo su consideración y quedando a disposición para presentar el protocolo y coordinar horarios de trabajo, saludo a usted atentamente.")
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(42)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Dr. Ruben Concha Cayola")
    set_run_font(r, bold=True)
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p2.add_run("Médico residente SAFCI")
    set_run_font(r)

    # Document properties and save.
    doc.core_properties.title = "Exposición ergonómica laboral y patrones de uso de AINEs en el CSI San Jorge"
    doc.core_properties.subject = "Avance para segunda revisión de monografía"
    doc.core_properties.author = "Ruben Concha Cayola"
    doc.core_properties.keywords = "ergonomía, AINEs, dolor osteoarticular, Zudáñez, atención primaria"
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
